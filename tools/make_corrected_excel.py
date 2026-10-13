# -*- coding: utf-8 -*-
"""Исправленные копии исходных Excel-файлов с подсветкой.

Берёт исходный файл из data/source, вносит в копию правки из реестра
data/corrections (ответы проверяющих), заново проверяет уже исправленный файл
и подсвечивает в нём то, что осталось:
  зелёная заливка + примечание — ячейка исправлена, в примечании «было»;
  жёлтая заливка + примечание — оставшаяся ошибка и вопрос к проверяющему.
Значения ячеек примечания не меняют, поэтому файл читается как прежде.
Исходные файлы не меняются. Копии кладутся в reports/ispravlennye под теми
же именами.

Запуск:
    python3 tools/make_corrected_excel.py                  # все файлы, по которым есть правки
    python3 tools/make_corrected_excel.py 2026-07_zapad_div1.xlsx ...
"""
import copy
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

import corrections
import report_data
from error_report import collect, period_sums, ERROR, SRC_DIR
from hockey_data import read_workbook, norm, parse_player, FIRST_COL

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "reports", "ispravlennye")
AUTHOR = "Проверка данных"
FIXED = PatternFill("solid", fgColor="FF92D050")     # исправлено
WRONG = PatternFill("solid", fgColor="FFFFFF00")     # осталась ошибка

SCORE_KINDS = {"счёт не сходится с периодами", "похоже на буллиты, но пометки нет",
               "итоговый счёт записан наоборот", "неправдоподобный счёт", "пустые периоды",
               "период не разобран", "итог не разобран", "пустой итоговый счёт"}


def note(cell, text, fill):
    """Заливка и примечание; если примечание уже есть — дописываем."""
    cell.fill = fill
    if cell.comment:
        text = cell.comment.text + "\n\n" + text
    c = Comment(text, AUTHOR)
    c.width = 360
    c.height = max(90, 18 * (len(text) // 45 + text.count("\n") + 2))
    cell.comment = c


def last_col(ws):
    return max(ws.max_column, 8)


def write_fixes(wb, applied):
    """Вносит применённые правки в ячейки книги."""
    for a in applied:
        g, fx = a["game"], a["fix"]
        ws = wb[g["sheet"]]
        fr = g["labels_seen"]["final"]
        pr = g["labels_seen"]["periods"]
        why = f"Исправлено по ответу проверяющего ({fx.get('note', '')})."
        if "final" in fx:
            c = ws.cell(fr, FIRST_COL)
            was = c.value
            c.value = fx["final"]
            note(c, f"{why}\nБыло: {was}", FIXED)
        if "periods" in fx:
            for i in range(FIRST_COL, last_col(ws) + 1):
                c = ws.cell(pr, i)
                k = i - FIRST_COL
                new = fx["periods"][k] if k < len(fx["periods"]) else None
                if norm(c.value) != (new or ""):
                    was = c.value
                    c.value = new
                    note(c, f"{why}\nБыло: {was if was not in (None, '') else 'пусто'}", FIXED)
        if fx.get("shootout"):
            # пометку ставим так же, как в остальных матчах файла: справа от итога
            c = ws.cell(fr, FIRST_COL + 1)
            if norm(c.value):
                raise RuntimeError(f"{g['sheet']}!{c.coordinate}: ячейка для пометки занята")
            c.value = corrections.SHOOTOUT_NOTE
            sample = template_note(ws)
            if sample is not None:
                c.font = copy.copy(sample.font)
                c.alignment = copy.copy(sample.alignment)
            note(c, f"{why}\nБыло: пометки не было", FIXED)


def template_note(ws):
    """Ячейка с пометкой «Б Буллиты», уже стоящая в файле, — образец оформления."""
    for row in ws.iter_rows(min_col=FIRST_COL + 1, max_col=FIRST_COL + 1):
        for c in row:
            if norm(c.value) == corrections.SHOOTOUT_NOTE:
                return c
    return None


def cells_for(ws, g, p):
    """Ячейки матча, к которым относится ошибка."""
    L = g["labels_seen"]
    what = p["what"]
    if what in SCORE_KINDS:
        out = [ws.cell(L["final"], FIRST_COL)]
        if "periods" in L:
            row = L["periods"]
            filled = [ws.cell(row, c) for c in range(FIRST_COL, last_col(ws) + 1) if norm(ws.cell(row, c).value)]
            out += filled or [ws.cell(row, FIRST_COL)]
        return out
    if what == "нет судьи" and "judge" in L:
        return [ws.cell(L["judge"] + 1, FIRST_COL)]
    if what == "нет секретаря" and "secretary" in L:
        return [ws.cell(L["secretary"] + 1, FIRST_COL)]
    if what == "матчи не из этого дивизиона":
        return [ws.cell(L[k], FIRST_COL) for k in ("comp", "team1", "team2") if k in L]
    if p.get("players"):
        out = []
        for key in ("roster1", "roster2"):
            if key not in L:
                continue
            row = L[key] + 1
            for c in range(FIRST_COL, last_col(ws) + 1):
                v = norm(ws.cell(row, c).value)
                if v and parse_player(v)[1] in p["players"]:
                    out.append(ws.cell(row, c))
        if out:
            return out
    return [ws.cell(g["row"], 1)]


def mark_errors(wb, problems, all_games):
    by_row = {}
    for g in all_games:
        by_row.setdefault(g["row"], []).append(g)
    marked = 0
    for p in problems:
        if p["kind"] != ERROR:
            continue
        text = "ОШИБКА: " + p["text"] + (f"\n{p['fix']}" if p.get("fix") else "")
        games = [p["game"]] if p["game"] is not None else [x for r in p["rows"] for x in by_row.get(r, [])]
        for g in games:
            ws = wb[g["sheet"]]
            extra = ""
            if p["what"] in SCORE_KINDS and g["periods_raw"]:
                sh, sa = period_sums(g)
                if sh is not None:
                    extra = (f"\nИтог {g['final_raw']}"
                             + (f" ({' '.join(g['final_notes'])})" if g["final_notes"] else "")
                             + f", по периодам {' '.join(g['periods_raw'])} = {sh}:{sa}.")
            for c in cells_for(ws, g, p):
                note(c, text.replace("\n", extra + "\n", 1) if p.get("fix") else text + extra, WRONG)
                marked += 1
    return marked


def build(fname):
    src = os.path.join(SRC_DIR, fname)
    dst = os.path.join(OUT_DIR, fname)
    games, _ = read_workbook(src)
    _, corr = corrections.apply(fname, games)
    if corr["conflicts"]:
        raise RuntimeError(f"{fname}: правки не сходятся с исходным файлом: {corr['conflicts']}")

    wb = openpyxl.load_workbook(src)
    write_fixes(wb, corr["applied"])
    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(dst)

    # Проверяем уже исправленный файл: правки из реестра должны оказаться «уже внесены».
    games2, problems, summary = collect(dst)
    if summary["applied"] or summary["conflicts"]:
        raise RuntimeError(f"{fname}: в исправленном файле правки внесены не все")
    wb = openpyxl.load_workbook(dst)
    marked = mark_errors(wb, problems, summary["all_games"])
    wb.save(dst)
    errors = sum(1 for p in problems if p["kind"] == ERROR)

    # Отчёт об оставшихся ошибках — рядом с файлом.
    base = os.path.splitext(fname)[0]
    js = os.path.join(OUT_DIR, f".{base}.json")
    report_data.main([dst], out=js, title="ОСТАВШИЕСЯ ОШИБКИ",
                     note=f"Исправленный файл {fname}: правки проверяющих внесены, "
                          f"ниже — что осталось исправить")
    docx = os.path.join(OUT_DIR, f"{base}_ostavshiesya_oshibki.docx")
    subprocess.run(["node", os.path.join(ROOT, "tools", "docx_build", "make_docx.js"), js, docx],
                   check=True, stdout=subprocess.DEVNULL)
    os.remove(js)
    return dict(file=fname, fixes=len(corr["applied"]), errors=errors, marked=marked,
                misplaced=len(corr["misplaced"]))


def main(args):
    files = args or sorted({fx["file"] for fx in corrections.FIXES})
    for f in files:
        r = build(os.path.basename(f))
        print(f"{r['file']}: внесено правок {r['fixes']}, осталось ошибок {r['errors']}, "
              f"подсвечено ячеек {r['marked']}")


if __name__ == "__main__":
    main(sys.argv[1:])
