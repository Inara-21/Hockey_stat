# -*- coding: utf-8 -*-
"""Полная проверка исходного файла на ошибки любого рода.

Ничего не исправляет и не формирует — только читает и сообщает.
Запуск:  python3 tools/check_file.py data/new/div1.xlsx
"""
import os, sys, datetime, difflib
from collections import Counter, defaultdict
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from hockey_data import read_workbook, parse_score, LABELS

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REQUIRED = ["comp", "date", "time", "team1", "team2", "roster1", "roster2",
            "secretary", "judge", "final", "periods"]
SEPS = ("*", ":", "-", "–", "—", "х", "x")


def sep_of(t):
    for s in SEPS:
        if s in (t or ""):
            return s
    return "нет"


def main(path):
    out, errors, warns, notes = [], [], [], []
    say = lambda t="": (out.append(t), print(t))[0]
    games, skipped = read_workbook(path)
    err = lambda g, t: errors.append((g["row"], t))
    warn = lambda g, t: warns.append((g["row"], t))

    say("=" * 78)
    say(f"ПРОВЕРКА: {os.path.basename(path)}")
    say("=" * 78)

    # 1 -------------------------------------------------- общее
    say("\n1. ОБЩЕЕ")
    say(f"   матчей: {len(games)}   листы: {dict(Counter(g['sheet'] for g in games))}")
    if skipped:
        say(f"   листы без матчей (пропущены): {skipped}")
    say(f"   соревнования: {dict(Counter(g['comp'] for g in games))}")
    say(f"   формат даты: {dict(Counter(g['date_kind'] for g in games))}"
        f"   времени: {dict(Counter(g['time_kind'] for g in games))}")
    say(f"   составы (к1,к2): {dict(sorted(Counter((len(g['roster1']), len(g['roster2'])) for g in games).items()))}")
    say(f"   периодов в матче: {dict(sorted(Counter(len(g['periods_raw']) for g in games).items()))}")

    # 2 -------------------------------------------------- структура блоков
    say("\n2. СТРУКТУРА БЛОКОВ")
    miss = Counter()
    for g in games:
        for f in REQUIRED:
            if f not in g["labels_seen"]:
                miss[f] += 1
                err(g, f"нет строки «{LABELS[f][0]}»")
    say("   все 11 полей на месте во всех блоках" if not miss else f"   отсутствуют: {dict(miss)}")
    swapped = sum(1 for g in games
                  if g["labels_seen"].get("team2", 9e9) < g["labels_seen"].get("team1", 0))
    say(f"   блоков с переставленными «Команда 1»/«Команда 2»: {swapped}"
        + ("  (читаем по подписям, на данные не влияет)" if swapped else ""))

    # 3 -------------------------------------------------- значения
    say("\n3. ЗНАЧЕНИЯ ПОЛЕЙ")
    for g in games:
        if not g["comp"]: err(g, "пустое соревнование")
        if g["date"] is None: err(g, f"дата не разобрана ({g['date_kind']})")
        if g["time"] is None: err(g, f"время не разобрано ({g['time_kind']})")
        if not g["team1"] or not g["team2"]: err(g, "пустое название команды")
        elif g["team1"] == g["team2"]: err(g, f"команда против себя: {g['team1']}")
        if not g["roster1"] or not g["roster2"]: err(g, "пустой состав")
        if not g["judge"]: warn(g, "нет судьи")
        if not g["secretary"]: warn(g, "нет секретаря")
        if not g["final_raw"]: err(g, "пустой итоговый счёт")
        elif g["goals1"] is None: err(g, f"итог не разобран: '{g['final_raw']}'")
        if not g["periods_raw"]: err(g, "пустой счёт по периодам")
        for p in g["periods_raw"]:
            if parse_score(p)[0] is None: err(g, f"период не разобран: '{p}'")
        for who, roster in (("к1", g["roster1"]), ("к2", g["roster2"])):
            nums = [n for n, _ in roster]
            for n, c in Counter(nums).items():
                if c > 1: err(g, f"повтор номера {who}: №{n}")
            for n, nm in roster:
                if not nm: err(g, f"игрок без ФИО {who}: №{n}")
                if not n: warn(g, f"игрок без номера {who}: {nm}")
    say(f"   ошибок: {len(errors)}   предупреждений: {len(warns)}")

    # 4 -------------------------------------------------- счёт
    say("\n4. СЧЁТ")
    say(f"   разделитель в итоге: {dict(Counter(sep_of(g['final_raw']) for g in games))}")
    say(f"   разделитель в периодах: {dict(Counter(sep_of(p) for g in games for p in g['periods_raw']))}")
    mixed = [g for g in games if g["periods_raw"]
             and sep_of(g["final_raw"]) != sep_of(g["periods_raw"][0])]
    if mixed:
        notes.append(f"в {len(mixed)} матчах разделитель итога отличается от периодов")
        say(f"   РАЗНЫЙ разделитель итога и периодов: {len(mixed)} матчей")

    mism, ot_like = [], []
    for g in games:
        if g["goals1"] is None or any(parse_score(p)[0] is None for p in g["periods_raw"]):
            continue
        sh = sum(parse_score(p)[0] for p in g["periods_raw"])
        sa = sum(parse_score(p)[1] for p in g["periods_raw"])
        if (sh, sa) != (g["goals1"], g["goals2"]):
            mism.append((g, sh, sa))
            if sh == sa and (g["goals1"] - sh) + (g["goals2"] - sa) == 1:
                ot_like.append(g)
            warn(g, f"сумма периодов {sh}:{sa} != итог {g['final_raw']}")
    say(f"   сумма периодов != итог: {len(mism)} матчей")
    say(f"      из них после 3 периодов ничья и +1 гол (похоже на овертайм): {len(ot_like)}")
    for g, sh, sa in mism:
        if g not in ot_like:
            say(f"      НЕ объясняется овертаймом: стр.{g['row']} {g['date']} {g['time']} "
                f"{g['team1']} — {g['team2']}: периоды {sh}:{sa}, итог {g['final_raw']}")
    gs = [(g["goals1"], g["goals2"]) for g in games if g["goals1"] is not None]
    if gs:
        say(f"   голов за матч: {min(a+b for a,b in gs)}..{max(a+b for a,b in gs)}, "
            f"в среднем {sum(a+b for a,b in gs)/len(gs):.1f}; ничьих {sum(1 for a,b in gs if a==b)}")
        if any(a < 0 or b < 0 for a, b in gs): say("   ОШИБКА: отрицательный счёт")

    # 5 -------------------------------------------------- повторы и пересечения
    say("\n5. ПОВТОРЫ И ПЕРЕСЕЧЕНИЯ")
    key = lambda g: (g["comp"], str(g["date"]), g["time"], g["team1"], g["team2"])
    dups = [k for k, c in Counter(key(g) for g in games).items() if c > 1]
    say(f"   одинаковых матчей: {len(dups)}")
    for d in dups[:10]: say(f"      {d}")
    clash = 0
    for (d, t), gs2 in sorted(defaultdict(list, {
            (str(g['date']), g['time']): [x for x in games
                                          if str(x['date']) == str(g['date']) and x['time'] == g['time']]
            for g in games}).items()):
        cnt = Counter()
        for g in gs2: cnt[g["team1"]] += 1; cnt[g["team2"]] += 1
        rep = [t2 for t2, c in cnt.items() if c > 1]
        if rep:
            clash += 1
            if clash <= 10: say(f"      {d} {t}: команда в двух матчах сразу — {rep}")
    say(f"   пересечений по времени: {clash}")

    # 6 -------------------------------------------------- команды и игроки
    say("\n6. КОМАНДЫ И ИГРОКИ")
    tm, ptm, pnum = Counter(), defaultdict(Counter), defaultdict(Counter)
    for g in games:
        tm[g["team1"]] += 1; tm[g["team2"]] += 1
        for team, roster in ((g["team1"], g["roster1"]), (g["team2"], g["roster2"])):
            for n, nm in roster:
                ptm[nm][team] += 1; pnum[nm][n] += 1
    say(f"   команд: {len(tm)}, матчей у каждой: {dict(tm)}")
    if tm: say(f"   разброс матчей между командами: {max(tm.values()) - min(tm.values())}")
    names = sorted(tm)
    for i, a in enumerate(names):
        for b in names[i+1:]:
            if difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio() > 0.8:
                notes.append(f"похожие названия команд: '{a}' и '{b}'")
    say(f"   уникальных игроков: {len(ptm)}")
    multi = {n: dict(t) for n, t in ptm.items() if len(t) > 1}
    say(f"   игроков в разных командах: {len(multi)}")
    for n, t in list(multi.items())[:10]: say(f"      {n}: {t}")
    mn = {n: dict(c) for n, c in pnum.items() if len(c) > 1}
    say(f"   игроков с разными номерами: {len(mn)}")
    for n, c in list(mn.items())[:10]: say(f"      {n}: {c}")
    pl = sorted(ptm)
    for i, a in enumerate(pl):
        for b in pl[i+1:]:
            if difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio() > 0.9:
                notes.append(f"похожие ФИО: '{a}' и '{b}'")
    gk = defaultdict(Counter)
    for g in games:
        for team, roster in ((g["team1"], g["roster1"]), (g["team2"], g["roster2"])):
            if roster: gk[team][roster[0][1]] += 1
    for t in sorted(gk):
        if len(gk[t]) > 1:
            notes.append(f"у «{t}» первым в составе (вратарём) указаны разные игроки: {dict(gk[t])}")
    say(f"   судьи: {dict(Counter(g['judge'] for g in games))}")
    say(f"   секретари: {dict(Counter(g['secretary'] for g in games))}")

    # 7 -------------------------------------------------- календарь
    say("\n7. КАЛЕНДАРЬ")
    byday = Counter(g["date"] for g in games if g["date"])
    if byday:
        days = sorted(byday)
        say(f"   период: {days[0]} .. {days[-1]}, дней с матчами: {len(days)}")
        gap = [days[0] + datetime.timedelta(d) for d in range((days[-1]-days[0]).days + 1)]
        empty = [d for d in gap if d not in byday]
        say(f"   дней без матчей внутри периода: {len(empty)}" + (f" — {empty[:10]}" if empty else ""))
        say(f"   матчей в день: {dict(sorted(Counter(byday.values()).items()))}")
        avg = sum(byday.values()) / len(byday)
        low = [(str(d), byday[d]) for d in days if byday[d] < avg / 2]
        if low:
            notes.append(f"дни с необычно малым числом матчей: {low}")
            say(f"   заметно меньше матчей, чем обычно: {low}")

    # 8 -------------------------------------------------- ячейки
    say("\n8. ДЕФЕКТЫ ЯЧЕЕК")
    wb = openpyxl.load_workbook(path)
    lead = trail = nbsp = dbl = form = 0
    merged = []
    for ws in wb.worksheets:
        merged += [str(r) for r in ws.merged_cells.ranges]
        for row in ws.iter_rows(max_row=ws.max_row, max_col=ws.max_column):
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v != v.lstrip(): lead += 1
                    if v != v.rstrip(): trail += 1
                    if "\xa0" in v: nbsp += 1
                    if "  " in v: dbl += 1
                    if v.startswith("="): form += 1
    say(f"   пробел в начале: {lead}, в конце: {trail}, неразрывный: {nbsp}, двойной: {dbl}, формул: {form}")
    say(f"   объединённых диапазонов: {len(merged)}" + (f" — {merged[:5]}" if merged else ""))
    hidden_r = sum(1 for ws in wb.worksheets for d in ws.row_dimensions.values() if d.hidden)
    hidden_c = sum(1 for ws in wb.worksheets for d in ws.column_dimensions.values() if d.hidden)
    say(f"   скрытых строк: {hidden_r}, скрытых столбцов: {hidden_c}")

    # ------------------------------------------------------ итог
    if errors:
        say(f"\n{'=' * 78}\nОШИБКИ ({len(errors)})")
        for row, t in errors[:80]: say(f"   стр.{row}: {t}")
    if warns:
        say(f"\n{'=' * 78}\nПРЕДУПРЕЖДЕНИЯ ({len(warns)})")
        for row, t in warns[:80]: say(f"   стр.{row}: {t}")
    if notes:
        say(f"\n{'=' * 78}\nЗАМЕЧАНИЯ ({len(notes)})")
        for t in dict.fromkeys(notes): say(f"   {t}")

    say(f"\n{'=' * 78}")
    say(f"ИТОГ: ошибок {len(errors)}, предупреждений {len(warns)}, замечаний {len(set(notes))}")
    say("=" * 78)

    os.makedirs(os.path.join(ROOT, "reports"), exist_ok=True)
    with open(os.path.join(ROOT, "reports", "proverka.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(out) + "\n")
    return len(errors)


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "data", "new", "div1.xlsx")
    sys.exit(1 if main(p) else 0)
