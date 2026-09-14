# -*- coding: utf-8 -*-
"""Генератор хоккейных протоколов игрового дня: data/games.xlsx -> reports/*.pdf

Принцип: данные читаются ПО ПОЗИЦИИ (жёсткий блок 15 строк на игру) и переносятся
в PDF без изменений. Подписи в левом столбце используются только для контроля
(в исходнике встречаются опечатки), но не влияют на разбор.

Запуск:  python3 tools/build_reports.py
"""
import openpyxl, html, subprocess, os, sys, shutil
from collections import defaultdict, Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "games.xlsx")
OUT = os.path.join(ROOT, "reports")
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

ROWS_PER_GAME = 15          # жёсткая сетка блока одной игры
CARDS_PER_SHEET = 4         # карточек игр на листе А4 (2 x 2)
MAX_PLAYERS = 7             # строк в таблице состава

# смещения строк внутри блока игры (0-based от первой строки блока)
# Строки ЗНАЧЕНИЙ (откуда читаем данные)
OFF_COMP, OFF_DATE, OFF_TIME, OFF_T1, OFF_T2 = 0, 1, 2, 3, 4
OFF_ROSTER1, OFF_ROSTER2 = 6, 8
OFF_SECRETARY, OFF_JUDGE = 10, 12
OFF_FINAL, OFF_PERIODS = 13, 14

# Строки ПОДПИСЕЙ в левом столбце. У составов, секретарей и судей подпись стоит
# на строку ВЫШЕ своих значений, поэтому смещения тут другие.
EXPECTED_LABELS = {
    0: "Название соревнования", 1: "Дата", 2: "Время начала матча",
    3: "Команда 1", 4: "Команда2",
    5: "Состав команды 1", 7: "Состав команды 2",
    9: "Секретари", 11: "Судьи",
    13: "Итоговый счёт матча", 14: "Счёт по периодам",
}

# ---------------------------------------------------------------- чтение
def raw(ws, r, c):
    return ws.cell(row=r, column=c).value

def txt(ws, r, c):
    v = raw(ws, r, c)
    return "" if v is None else str(v).strip()

def parse_player(s):
    """'Мажоров Владислав Владимирович №1' -> ('1', 'Мажоров Владислав Владимирович').
    Терпимо к слитному написанию 'Игоревич№15'."""
    if not s:
        return ("", "")
    if "№" in s:
        name, num = s.rsplit("№", 1)
        return (num.strip(), name.strip())
    return ("", s.strip())

def parse_score(s):
    """'3*5' -> (3, 5, '3 : 5'). Числа не меняются, меняется только разделитель."""
    s = (s or "").strip()
    for sep in ("*", ":", "-"):
        if sep in s:
            a, b = s.split(sep, 1)
            a, b = a.strip(), b.strip()
            try:
                return (int(a), int(b), f"{a} : {b}")
            except ValueError:
                return (None, None, s)
    return (None, None, s)

def fmt_date(v):
    try:
        return v.strftime("%d.%m.%Y")
    except Exception:
        return str(v)[:10]

def fmt_time(v):
    try:
        return v.strftime("%H:%M")
    except Exception:
        return str(v)[:5]

def read_games(path):
    """Читает все игры. Возвращает (games, label_warnings)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    games, label_warnings = [], []
    for ws in wb.worksheets:
        if ws.max_row % ROWS_PER_GAME != 0:
            raise SystemExit(
                f"ОШИБКА: лист '{ws.title}': {ws.max_row} строк не делится на {ROWS_PER_GAME}. "
                "Структура блоков нарушена — разбор остановлен.")
        for i in range(ws.max_row // ROWS_PER_GAME):
            base = i * ROWS_PER_GAME + 1
            # контроль подписей (не влияет на разбор, только предупреждения)
            for off, expected in EXPECTED_LABELS.items():
                got = txt(ws, base + off, 1)
                if got != expected:
                    label_warnings.append(
                        dict(sheet=ws.title, game=i + 1, row=base + off,
                             expected=expected, got=got))
            final_raw = txt(ws, base + OFF_FINAL, 2)
            fh, fa, final_disp = parse_score(final_raw)
            t1, t2 = txt(ws, base + OFF_T1, 2), txt(ws, base + OFF_T2, 2)
            periods_raw = [txt(ws, base + OFF_PERIODS, c)
                           for c in range(2, 9) if txt(ws, base + OFF_PERIODS, c)]
            reg, ot = periods_raw[:3], periods_raw[3:]
            if fh is None or fa is None:
                winner = ""
            elif fh > fa:
                winner = t1
            elif fa > fh:
                winner = t2
            else:
                winner = "ничья"
            date_v = raw(ws, base + OFF_DATE, 2)
            games.append(dict(
                sheet=ws.title, index=i + 1, row=base,
                comp=txt(ws, base + OFF_COMP, 2),
                date=fmt_date(date_v), date_key=str(date_v)[:10],
                time=fmt_time(raw(ws, base + OFF_TIME, 2)),
                t1=t1, t2=t2,
                roster1=[parse_player(txt(ws, base + OFF_ROSTER1, c))
                         for c in range(2, 9) if txt(ws, base + OFF_ROSTER1, c)],
                roster2=[parse_player(txt(ws, base + OFF_ROSTER2, c))
                         for c in range(2, 9) if txt(ws, base + OFF_ROSTER2, c)],
                secretary=txt(ws, base + OFF_SECRETARY, 2),
                judge=txt(ws, base + OFF_JUDGE, 2),
                final_raw=final_raw, final=final_disp, fh=fh, fa=fa, winner=winner,
                periods_raw=periods_raw,
                reg=", ".join(parse_score(p)[2] for p in reg),
                ot=", ".join(parse_score(p)[2] for p in ot),
            ))
    return games, label_warnings

# ---------------------------------------------------------------- вёрстка
CSS = '''
* { box-sizing: border-box; }
@page { size: A4; margin: 21mm 14mm 14mm 14mm; }
body { font-family:"DejaVu Sans","Liberation Sans",sans-serif; color:#1a1f36; margin:0; font-size:9.5px; }
.sheet { page-break-after: always; }
.sheet:last-child { page-break-after: auto; }
.head { display:flex; align-items:center; justify-content:space-between; margin-bottom:8mm; }
.head .logo { font-weight:800; font-size:16px; letter-spacing:1px; color:#2b2f77; }
.head .title { font-weight:800; font-size:17px; letter-spacing:1px; }
.head .date { font-weight:800; font-size:15px; }
.grid { display:grid; grid-template-columns:1fr 1fr; gap:7mm; }
.card { border:1.3px solid #3a3f6b; border-radius:9px; padding:4.5mm 5mm; }
.fld { display:flex; align-items:baseline; gap:5px; border-bottom:1px solid #b9bdd6;
       padding:2px 2px; min-height:14px; }
.fld.wide { margin-bottom:3mm; }
.fld.wide2 { margin-top:2.2mm; }
.row2 { display:grid; grid-template-columns:1fr 1fr; gap:5mm; margin-bottom:2.2mm; }
.lb { font-weight:800; font-size:8px; letter-spacing:.4px; white-space:nowrap; color:#2b2f4c; }
.vl { font-style:italic; color:#1a1f36; }
.rosters { display:grid; grid-template-columns:1fr 1fr; gap:5mm; margin:2.6mm 0; }
table.ros { width:100%; border-collapse:collapse; }
table.ros th, table.ros td { border:1px solid #b9bdd6; padding:1.8px 4px; font-size:8.3px; }
table.ros th { font-weight:800; font-size:7.5px; background:#f2f3f9; letter-spacing:.3px; }
table.ros td.num, table.ros th.num { width:22px; text-align:center; }
table.ros td.pl { font-style:italic; }
.officials { border:1.3px solid #3a3f6b; border-radius:9px; padding:5mm 6mm; margin-top:6mm;
             display:grid; grid-template-columns:1fr 1fr; gap:10mm; break-inside:avoid; }
.off-h { font-weight:800; font-size:9px; letter-spacing:.5px; margin-bottom:3.5mm; }
.off-v { font-style:italic; border-bottom:1px solid #b9bdd6; padding-bottom:3px; }
'''

def e(x):
    return html.escape(str(x))

def roster_table(roster):
    rows = ""
    for i in range(MAX_PLAYERS):
        num, name = roster[i] if i < len(roster) else ("", "")
        rows += f'<tr><td class="num">{e(num)}</td><td class="pl">{e(name)}</td></tr>'
    return ('<table class="ros"><tr><th class="num">№</th><th class="pl">ИГРОК</th></tr>'
            + rows + '</table>')

def card_html(g):
    return f'''<div class="card">
  <div class="fld wide"><span class="lb">СОРЕВНОВАНИЕ</span><span class="vl">{e(g["comp"])}</span></div>
  <div class="row2">
    <div class="fld"><span class="lb">ДАТА</span><span class="vl">{e(g["date"])}</span></div>
    <div class="fld"><span class="lb">ВРЕМЯ</span><span class="vl">{e(g["time"])}</span></div>
  </div>
  <div class="row2">
    <div class="fld"><span class="lb">КОМАНДА&nbsp;А</span><span class="vl">{e(g["t1"])}</span></div>
    <div class="fld"><span class="lb">КОМАНДА&nbsp;В</span><span class="vl">{e(g["t2"])}</span></div>
  </div>
  <div class="rosters">{roster_table(g["roster1"])}{roster_table(g["roster2"])}</div>
  <div class="row2">
    <div class="fld"><span class="lb">ИТОГОВЫЙ&nbsp;СЧЁТ</span><span class="vl">{e(g["final"])}</span></div>
    <div class="fld"><span class="lb">ПОБЕДИТЕЛЬ</span><span class="vl">{e(g["winner"])}</span></div>
  </div>
  <div class="fld wide2"><span class="lb">СЧЁТ&nbsp;ПО&nbsp;ПЕРИОДАМ</span><span class="vl">{e(g["reg"])}</span></div>
  <div class="fld wide2"><span class="lb">ОВЕРТАЙМЫ</span><span class="vl">{e(g["ot"])}</span></div>
</div>'''

def officials_html(judge, secretary):
    return f'''<div class="officials">
  <div class="off-col"><div class="off-h">СУДЬИ</div><div class="off-v">{e(judge)}</div></div>
  <div class="off-col"><div class="off-h">СЕКРЕТАРИ</div><div class="off-v">{e(secretary)}</div></div>
</div>'''

def day_sheets_html(day_games):
    """Листы одного игрового дня. Судьи/секретари — внизу последнего листа."""
    date = day_games[0]["date"]
    judge, secretary = day_games[0]["judge"], day_games[0]["secretary"]
    chunks = [day_games[i:i + CARDS_PER_SHEET]
              for i in range(0, len(day_games), CARDS_PER_SHEET)]
    out = ""
    for si, chunk in enumerate(chunks):
        label = e(date) + (f' · лист {si+1} из {len(chunks)}' if len(chunks) > 1 else "")
        off = officials_html(judge, secretary) if si == len(chunks) - 1 else ""
        out += (f'<div class="sheet"><div class="head">'
                f'<div class="logo">ХОККЕЙ</div>'
                f'<div class="title">ПРОТОКОЛ ИГРОВОГО ДНЯ</div>'
                f'<div class="date">{label}</div></div>'
                f'<div class="grid">{"".join(card_html(g) for g in chunk)}</div>{off}</div>')
    return out, len(chunks)

def page(body):
    return ('<!doctype html><html><head><meta charset="utf-8">'
            f'<style>{CSS}</style></head><body>{body}</body></html>')

def render_pdf(html_str, pdf_path, tmp_html):
    with open(tmp_html, "w", encoding="utf-8") as f:
        f.write(html_str)
    subprocess.run([CHROME, "--headless", "--no-sandbox", "--disable-gpu",
                    "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}",
                    f"file://{tmp_html}"], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

# ---------------------------------------------------------------- main
def main():
    games, label_warnings = read_games(XLSX)
    print(f"Прочитано игр: {len(games)}")

    byday = defaultdict(list)
    for g in games:
        byday[(g["sheet"], g["date_key"])].append(g)
    days = sorted(byday.keys(), key=lambda k: (k[1], k[0]))

    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    os.makedirs(os.path.join(OUT, "by_day"), exist_ok=True)
    tmp = os.path.join(OUT, "_tmp.html")

    # Копим листы по месяцам: имя листа Excel ("июль", "август") = месяц.
    months = defaultdict(str)
    month_sheets = Counter()
    month_order = []
    total_sheets = 0
    for key in days:
        sheet_name, date_key = key
        day_games = byday[key]
        body, nsheets = day_sheets_html(day_games)
        total_sheets += nsheets
        if sheet_name not in months:
            month_order.append(sheet_name)
        months[sheet_name] += body
        month_sheets[sheet_name] += nsheets
        render_pdf(page(body),
                   os.path.join(OUT, "by_day", f"{date_key}_{sheet_name}.pdf"), tmp)

    for sheet_name in month_order:
        out_path = os.path.join(OUT, f"Протоколы_{sheet_name}.pdf")
        render_pdf(page(months[sheet_name]), out_path, tmp)
        n_games = sum(len(byday[k]) for k in days if k[0] == sheet_name)
        print(f"  {sheet_name}: игр {n_games}, листов {month_sheets[sheet_name]} -> {os.path.basename(out_path)}")

    os.remove(tmp)
    print(f"Дней: {len(days)} | листов всего: {total_sheets} | месяцев: {len(month_order)}")
    return games, label_warnings, days, byday, total_sheets

if __name__ == "__main__":
    main()
